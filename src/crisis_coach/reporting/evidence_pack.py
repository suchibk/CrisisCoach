"""Local DOCX export using verified file bytes and explicitly labelled gaps."""
import hashlib
import os
import tempfile
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from pydantic import BaseModel, ConfigDict
from ..models import SceneState
from ..models.evidence import EvidenceDefinition, EvidenceStatus
from ..models.reports import EvidencePackReference
from ..persistence.attachments import AttachmentStore
from ..tools.contracts import ToolError
from .completeness import calculate_completeness

class EvidencePackArgs(BaseModel):
    model_config = ConfigDict(extra="forbid")
    scene: SceneState
    items: tuple[EvidenceDefinition, ...]


def build_evidence_pack(args: EvidencePackArgs, attachment_root: Path, export_root: Path, event_id: str) -> EvidencePackReference:
    try:
        from docx import Document
        from docx.shared import Inches
        from docx.image.exceptions import UnrecognizedImageError, UnexpectedEndOfFileError, InvalidImageStreamError
    except ImportError as exc:
        raise ToolError('Report support is not installed; run: pip install -e ".[reports]"') from exc
    scene = args.scene
    store = AttachmentStore(attachment_root)
    document = Document()
    document.add_heading("Crisis Coach evidence pack", 0)
    generated = datetime.now(timezone.utc)
    document.add_paragraph(f"Incident: {scene.incident_id} | Person: {scene.person_name}")
    document.add_paragraph(f"Domain: {scene.domain_id} | Pack: {scene.pack_version} | Record revision: {scene.revision}")
    document.add_paragraph(f"Export generated at {generated.isoformat()} (not the collision time).")
    document.add_paragraph("This is a local evidence record for review, not a filed claim. Any AI image-quality assessment is not human verification or confirmation of claim readiness.")
    document.add_paragraph("Collision timestamp, location, and weather are not captured by this prototype and are not inferred. Any recorded source passages appear below; their legal applicability has not been determined.")
    # Reserve summary near the front; populate after validating and embedding files.
    summary = document.add_paragraph()
    table = document.add_table(rows=1, cols=4)
    for cell, label in zip(table.rows[0].cells, ("Required item", "Recorded status", "Stored in this pack", "Gap / issue")):
        cell.text = label
    labels = {item.item_id: item.label for item in args.items}
    problems = {}
    ordered = list(labels) + sorted(key for key, record in scene.evidence.items() if key not in labels and (record.attachments or record.note))
    for item_id in ordered:
        record = scene.evidence.get(item_id)
        if record is None: continue
        document.add_heading(labels.get(item_id, item_id + " (additional supplied evidence)"), 1)
        document.add_paragraph(f"Recorded status: {record.status.value}")
        if record.note:
            document.add_paragraph("Record note / user report: " + record.note)
        if not record.attachments:
            document.add_paragraph("No stored file accompanies this item.")
        for reference in record.attachments:
            document.add_paragraph(f"Attachment SHA-256: {reference.sha256} | {reference.size_bytes} bytes | {reference.media_type}")
            try:
                data = store.read(scene.incident_id, reference)
                if reference.media_type == "text/plain":
                    text = data.decode("utf-8")
                    document.add_paragraph("Original supplied text:")
                    document.add_paragraph(text)
                else:
                    document.add_picture(BytesIO(data), width=Inches(5.5))
                    document.add_paragraph("Stored image; see the evidence note for any AI quality assessment. No human verification is recorded.")
            except (ToolError, UnicodeError, ValueError, UnrecognizedImageError, UnexpectedEndOfFileError, InvalidImageStreamError) as exc:
                # Failed media must not prevent an honest partial report.
                problems[item_id] = "One or more attachments could not be included: " + type(exc).__name__
                document.add_paragraph(problems[item_id])
    completeness = calculate_completeness(args.items, scene.evidence, problems)
    summary.text = f"Stored applicable items: {completeness.stored_count}/{completeness.required_count} ({completeness.stored_percent}%). Gaps: {completeness.missing_count}. Reported-only items: {completeness.reported_only_count}. This is catalogue coverage, not an insurer completeness score."
    for row in completeness.items:
        cells = table.add_row().cells
        for cell, value in zip(cells, (row.label, row.status.value, "Yes" if row.stored else "No", row.problem or "")):
            cell.text = value
    if scene.knowledge_answers:
        document.add_heading("Source lookups recorded during the incident", 1)
        for answer in scene.knowledge_answers:
            document.add_paragraph("Question: " + answer.question)
            document.add_paragraph(answer.message)
            for citation in answer.citations:
                label = "SYNTHETIC DEMO — " if citation.synthetic else ""
                document.add_paragraph(f"{label}{citation.title}, section {citation.section} [{citation.source_id}]")
                document.add_paragraph(citation.text)
                document.add_paragraph("Source: " + citation.source_uri)
    if scene.last_decision:
        document.add_heading("Most recent collection ordering", 1)
        for candidate in scene.last_decision.candidates:
            document.add_paragraph(f"{candidate.item_id}: score {candidate.score:.4f}; {candidate.reason}")
    export_root = Path(export_root).resolve()
    incident_key = hashlib.sha256(scene.incident_id.encode()).hexdigest()
    # Each request gets a separate export; repeated requests cannot overwrite an older report.
    export_key = hashlib.sha256(event_id.encode()).hexdigest()
    relative = Path(incident_key) / f"evidence-{scene.revision}-{export_key}.docx"
    target = (export_root / relative).resolve()
    if not target.is_relative_to(export_root): raise ToolError("Invalid export path")
    target.parent.mkdir(parents=True, exist_ok=True)
    descriptor, filename = tempfile.mkstemp(dir=target.parent)
    temporary = Path(filename)
    try:
        with os.fdopen(descriptor, "wb") as stream:
            document.save(stream)
            stream.flush()
            os.fsync(stream.fileno())
        data = temporary.read_bytes()
        os.replace(temporary, target)
    finally:
        temporary.unlink(missing_ok=True)
    return EvidencePackReference(relative_path=relative.as_posix(), sha256=hashlib.sha256(data).hexdigest(),
        generated_at=generated, incident_revision=scene.revision, stored_items=completeness.stored_count,
        required_items=completeness.required_count, problems=tuple(f"{key}: {value}" for key,value in problems.items()))


def read_export(export_root: Path, incident_id: str, reference: EvidencePackReference) -> bytes:
    root = Path(export_root).resolve()
    incident_key = hashlib.sha256(incident_id.encode()).hexdigest()
    path = (root / reference.relative_path).resolve()
    if not path.is_relative_to(root / incident_key):
        raise ToolError("Export reference is outside this incident")
    try:
        data = path.read_bytes()
    except OSError as exc:
        raise ToolError("Export is missing or unreadable") from exc
    if hashlib.sha256(data).hexdigest() != reference.sha256:
        raise ToolError("Export failed integrity validation")
    return data
