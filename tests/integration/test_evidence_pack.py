import base64
import zipfile
from pathlib import Path
import pytest
from crisis_coach import CollisionWorkflow, SessionStatus
from crisis_coach.models.events import ReportEvent
from crisis_coach.persistence.sqlite import SQLiteIncidentRepository
from crisis_coach.tools.capture import build_capture_executor
from crisis_coach.tools.contracts import ToolError
from crisis_coach.reporting.evidence_pack import read_export

pytest.importorskip("docx")
from docx import Document

PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aOZsAAAAASUVORK5CYII=")

@pytest.fixture
def scene(tmp_path):
    repo = SQLiteIncidentRepository(tmp_path / "incidents.db")
    coach = CollisionWorkflow(repository=repo, executor=build_capture_executor(tmp_path / "attachments"))
    state, _ = coach.start()
    for answer in ("No", "No", "Yes"): coach.turn(state, answer)
    return coach, state, repo, tmp_path


def report_text(path):
    doc = Document(path)
    return "\n".join([p.text for p in doc.paragraphs] + [c.text for t in doc.tables for r in t.rows for c in r.cells])


def test_pack_embeds_images_and_original_text_and_labels_reports(scene):
    coach, state, repo, root = scene
    photo = root / "photo.png"
    photo.write_bytes(PNG)
    coach.turn(state, f"/attach wide_scene_photo {photo}")
    coach.turn(state, "/statement statement I was stationary.\nThe other vehicle hit mine.")
    coach.turn(state, "/collected other_plate ABC123")
    coach.turn(state, "/unavailable other_insurance Driver left")
    coach.turn(state, "/export")
    ref = state.exports[-1]
    path = root / "exports" / ref.relative_path
    text = report_text(path)
    assert "I was stationary.\nThe other vehicle hit mine." in text
    assert "ABC123" in text and "Driver left" in text
    assert "not an insurer completeness score" in text
    assert ref.stored_items == 2
    assert repo.load(state.incident_id).exports[-1] == ref
    with zipfile.ZipFile(path) as archive:
        assert any(name.startswith("word/media/") for name in archive.namelist())
    assert read_export(root / "exports", state.incident_id, ref)


def test_missing_attachment_generates_honest_partial_pack(scene):
    coach, state, _, root = scene
    coach.turn(state, "/statement statement I was stationary")
    record = state.evidence["statement"]
    (root / "attachments" / record.attachments[0].relative_path).unlink()
    coach.turn(state, "/export")
    ref = state.exports[-1]
    assert ref.stored_items == 0
    assert ref.problems
    assert "could not be included" in report_text(root / "exports" / ref.relative_path)


def test_export_before_capture_is_partial_and_does_not_complete_scene(scene):
    coach, state, _, root = scene
    instruction = state.last_instruction
    deadline = state.timer_deadline
    coach.turn(state, "/export")
    assert state.exports[-1].stored_items == 0
    assert state.status is SessionStatus.ACTIVE
    assert state.last_instruction == instruction
    assert state.timer_deadline == deadline


def test_stand_down_blocks_export(scene):
    coach, state, _, root = scene
    coach.turn(state, "My chest hurts")
    coach.turn(state, "/export")
    assert not state.exports
    assert not (root / "exports").exists()


def test_duplicate_export_event_is_not_rebuilt(scene):
    coach, state, _, root = scene
    event = ReportEvent()
    coach.handle_event(state, event)
    assert coach.handle_event(state, event) is None
    assert len(state.exports) == 1
    assert len(list((root / "exports").rglob("*.docx"))) == 1


def test_tampered_export_and_cross_incident_reference_are_rejected(scene):
    coach, state, _, root = scene
    coach.turn(state, "/export")
    ref = state.exports[-1]
    with pytest.raises(ToolError): read_export(root / "exports", "different-incident", ref)
    (root / "exports" / ref.relative_path).write_bytes(b"changed")
    with pytest.raises(ToolError): read_export(root / "exports", state.incident_id, ref)


def test_langgraph_report_node(scene):
    pytest.importorskip("langgraph")
    _, _, repo, root = scene
    coach = CollisionWorkflow(repository=repo, backend="langgraph", executor=build_capture_executor(root / "attachments"))
    state, _ = coach.start()
    for answer in ("No", "No", "Yes"): coach.turn(state, answer)
    coach.turn(state, "/export")
    assert state.exports
    assert coach.last_route == ("safety", "controls", "gate", "report")


def test_inapplicable_captured_items_are_retained_as_additional_evidence(scene):
    coach, state, _, root = scene
    photo = root / "photo.png"
    photo.write_bytes(PNG)
    coach.turn(state, f"/attach wide_scene_photo {photo}")
    coach.turn(state, "/scenario scuff")
    coach.turn(state, "/export")
    ref = state.exports[-1]
    assert ref.required_items == 3
    assert ref.stored_items == 0
    assert "additional supplied evidence" in report_text(root / "exports" / ref.relative_path)
